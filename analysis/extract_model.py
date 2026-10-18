#!/usr/bin/env python3
"""
Extract all figures, formulas, scenarios, sensitivity grids, Monte Carlo outputs,
capex modes, menu items, dayparts, staffing, opex, sources, and completeness flags
from the CORSA Jakarta Financial Model v0.6 (.xlsx).

Generates:
  - analysis/outputs/model_summary.json
  - analysis/outputs/model_figures_table.md
"""

import argparse
import json
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parent.parent
DEFAULT_MODEL = ROOT / "data" / "raw" / "2026-10-18-corsa-jakarta-financial-model-v0-6.xlsx"
OUT_DIR = ROOT / "analysis" / "outputs"


def clean_val(v):
    if v is None:
        return None
    if isinstance(v, float):
        if abs(v - round(v)) < 1e-9:
            return int(round(v))
        return round(v, 6)
    return v


def fmt_idr(v):
    if v is None or v == "":
        return "—"
    if isinstance(v, (int, float)):
        return f"Rp {v:,.0f}"
    return str(v)


def fmt_chf(v):
    if v is None or v == "":
        return "—"
    if isinstance(v, (int, float)):
        return f"CHF {v:,.0f}"
    return str(v)


def fmt_pct(v, decimals=1):
    if v is None or v == "":
        return "—"
    if isinstance(v, (int, float)):
        return f"{v * 100:.{decimals}f}%"
    return str(v)


def extract_model(model_path: Path):
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    wb_val = openpyxl.load_workbook(model_path, data_only=True)
    wb_fml = openpyxl.load_workbook(model_path, data_only=False)

    summary = {
        "model_file": str(model_path.relative_to(ROOT)) if model_path.is_relative_to(ROOT) else str(model_path),
        "sheet_names": wb_val.sheetnames,
        "sheet_count": len(wb_val.sheetnames),
    }

    # 1. Inputs & Switches (Inputs!A1:L103)
    ws_in = wb_val["Inputs"]
    ws_in_f = wb_fml["Inputs"]
    inputs_table = []
    for r in range(10, 104):
        key = clean_val(ws_in.cell(row=r, column=1).value)
        param = clean_val(ws_in.cell(row=r, column=3).value)
        if key and param:
            inputs_table.append({
                "row": r,
                "key": key,
                "category": clean_val(ws_in.cell(row=r, column=2).value),
                "parameter": param,
                "unit": clean_val(ws_in.cell(row=r, column=4).value),
                "conservative": {"cell": f"Inputs!E{r}", "value": str(ws_in.cell(row=r, column=5).value)[:10] if r == 11 else clean_val(ws_in.cell(row=r, column=5).value)},
                "base": {"cell": f"Inputs!F{r}", "value": str(ws_in.cell(row=r, column=6).value)[:10] if r == 11 else clean_val(ws_in.cell(row=r, column=6).value)},
                "optimistic": {"cell": f"Inputs!G{r}", "value": str(ws_in.cell(row=r, column=7).value)[:10] if r == 11 else clean_val(ws_in.cell(row=r, column=7).value)},
                "active": {"cell": f"Inputs!H{r}", "value": str(ws_in.cell(row=r, column=8).value)[:10] if r == 11 else clean_val(ws_in.cell(row=r, column=8).value)},
                "source_ids": clean_val(ws_in.cell(row=r, column=9).value),
                "derivation": clean_val(ws_in.cell(row=r, column=11).value),
                "evidence_status": clean_val(ws_in.cell(row=r, column=12).value),
            })

    summary["switches_and_inputs"] = {
        "active_scenario_code": {"cell": "Inputs!D4", "value": clean_val(ws_in["D4"].value)},
        "active_scenario_name": {"cell": "Inputs!E4", "value": clean_val(ws_in["E4"].value), "formula": ws_in_f["E4"].value},
        "capex_financing_mode_code": {"cell": "Inputs!H4", "value": clean_val(ws_in["H4"].value)},
        "capex_financing_mode_name": {"cell": "Inputs!I4", "value": clean_val(ws_in["I4"].value), "formula": ws_in_f["I4"].value},
        "parameters": inputs_table,
    }

    # 2. Completeness & Model Horizon Audit (PnL_Monthly, Annual, Dashboard, Research_Tracker)
    ws_pnl = wb_val["PnL_Monthly"]
    ws_pnl_f = wb_fml["PnL_Monthly"]
    populated_months = []
    empty_months = []
    for col_idx in range(3, 63):  # C (3) to BJ (62) = 60 months
        col_letter = openpyxl.utils.get_column_letter(col_idx)
        m_num = ws_pnl.cell(row=4, column=col_idx).value
        m_date = ws_pnl.cell(row=5, column=col_idx).value
        rev_val = ws_pnl.cell(row=14, column=col_idx).value
        rev_fml = ws_pnl_f.cell(row=14, column=col_idx).value
        if rev_val is not None or rev_fml is not None:
            populated_months.append({
                "col": col_letter,
                "month_index": clean_val(m_num),
                "operating_year": clean_val(ws_pnl.cell(row=6, column=col_idx).value),
                "date": str(m_date)[:10] if m_date else None,
                "ramp_up_pct": clean_val(ws_pnl.cell(row=7, column=col_idx).value),
                "seasonality_idx": clean_val(ws_pnl.cell(row=8, column=col_idx).value),
                "total_revenue_idr": clean_val(rev_val),
                "ebitda_idr": clean_val(ws_pnl.cell(row=36, column=col_idx).value),
                "net_income_idr": clean_val(ws_pnl.cell(row=40, column=col_idx).value),
                "closing_cash_idr": clean_val(ws_pnl.cell(row=48, column=col_idx).value),
                "project_fcf_idr": clean_val(ws_pnl.cell(row=51, column=col_idx).value),
            })
        else:
            empty_months.append({"col": col_letter, "month_index": col_idx - 2})

    ws_ann = wb_val["Annual"]
    ws_ann_f = wb_fml["Annual"]
    annual_rows = []
    for r in [5, 6, 7, 8, 9, 10, 11, 12, 13, 14, 15, 16, 17, 18, 19, 21, 22, 23, 24, 25, 28, 29, 30, 31, 32]:
        label = clean_val(ws_ann.cell(row=r, column=1).value)
        if label:
            annual_rows.append({
                "row": r,
                "label": label,
                "Y1": {"cell": f"Annual!C{r}", "value": clean_val(ws_ann.cell(row=r, column=3).value), "formula": ws_ann_f.cell(row=r, column=3).value},
                "Y2": {"cell": f"Annual!D{r}", "value": clean_val(ws_ann.cell(row=r, column=4).value), "formula": ws_ann_f.cell(row=r, column=4).value},
                "Y3": {"cell": f"Annual!E{r}", "value": clean_val(ws_ann.cell(row=r, column=5).value), "formula": ws_ann_f.cell(row=r, column=5).value},
                "Y4": {"cell": f"Annual!F{r}", "value": clean_val(ws_ann.cell(row=r, column=6).value), "formula": ws_ann_f.cell(row=r, column=6).value},
                "Y5": {"cell": f"Annual!G{r}", "value": clean_val(ws_ann.cell(row=r, column=7).value), "formula": ws_ann_f.cell(row=r, column=7).value},
                "Total": {"cell": f"Annual!H{r}", "value": clean_val(ws_ann.cell(row=r, column=8).value)},
            })

    ws_rt = wb_val["Research_Tracker"]
    research_tracker = []
    rt_counts = {}
    for r in range(5, 18):
        q = clean_val(ws_rt.cell(row=r, column=1).value)
        if q:
            status = clean_val(ws_rt.cell(row=r, column=3).value)
            rt_counts[status] = rt_counts.get(status, 0) + 1
            research_tracker.append({
                "row": r,
                "cell_range": f"Research_Tracker!A{r}:F{r}",
                "question": q,
                "where_in_model": clean_val(ws_rt.cell(row=r, column=2).value),
                "status": status,
                "next_research_step": clean_val(ws_rt.cell(row=r, column=4).value),
                "who_how": clean_val(ws_rt.cell(row=r, column=5).value),
                "sources_found_so_far": clean_val(ws_rt.cell(row=r, column=6).value),
            })

    summary["completeness_audit"] = {
        "pnl_monthly_populated_count": len(populated_months),
        "pnl_monthly_populated_range": "PnL_Monthly!C4:O53 (Months 1 to 13, Jul 2027 to Jul 2028)",
        "pnl_monthly_empty_count": len(empty_months),
        "pnl_monthly_empty_range": "PnL_Monthly!P4:BJ53 (Months 14 to 60 completely blank)",
        "populated_months": populated_months,
        "annual_rows": annual_rows,
        "consequences_of_truncation": [
            "Annual!D5 (Year 2 revenue = Rp 765,055,200 / CHF 35,419) uses SUMIFS over PnL_Monthly operating year 2, where only column O (Month 13) is populated and columns P:Z (Months 14-24) are blank.",
            "Annual!E5:G5 (Years 3-5 revenue) and Annual!E11:G11 (Years 3-5 EBITDA) evaluate to 0 because PnL_Monthly columns AA:BJ (Months 25-60) are blank.",
            "Dashboard!B17 / C17 (Year-3 EBITDA = Rp 0 / CHF 0) references Annual!E11, which is 0 due to missing Months 25-36.",
            "Dashboard!B19 / C19 & Annual!C31:D31 (5-yr NPV at 15% discount rate = Rp -3,185,886,005 / CHF -147,495) discounts only 13 months of project free cash flow (+Rp 75.55m deposit return in Y5) against full Year-0 investment (Rp -4,266,246,130).",
            "Dashboard!B25 (Payback month = '> 60 months') checks PnL_Monthly!C53:BJ53 where cumulative FCF less investment (PnL_Monthly!O52 = Rp -3,052,259,680) only runs through Month 13.",
            "Dashboard!B26 & Annual!C30 (5-yr IRR = -46.9%) uses IRR(Annual!C32:H32) with cash flows [-4.266B, +1.101B, +113.2M, 0, 0, +75.5M].",
            "Dashboard!B18 / C18 (Lowest cash balance over 60 months = Rp 794,813,727 / CHF 36,797) uses MIN(PnL_Monthly!C48:BJ48), which ignores blank cells P48:BJ48 and reports Month 1 (PnL_Monthly!C48) across Months 1-13 only."
        ],
        "research_tracker_counts": rt_counts,
        "research_tracker_rows": research_tracker,
    }

    # 3. Dashboard outputs (Dashboard!A5:E35)
    ws_dash = wb_val["Dashboard"]
    ws_dash_f = wb_fml["Dashboard"]
    dashboard_kpis = []
    for r in range(8, 30):
        metric = clean_val(ws_dash.cell(row=r, column=1).value)
        if metric:
            dashboard_kpis.append({
                "row": r,
                "metric": metric,
                "col_b_cell": f"Dashboard!B{r}",
                "col_b_val": clean_val(ws_dash.cell(row=r, column=2).value),
                "col_b_formula": ws_dash_f.cell(row=r, column=2).value,
                "col_c_cell": f"Dashboard!C{r}",
                "col_c_val": clean_val(ws_dash.cell(row=r, column=3).value),
                "col_c_formula": ws_dash_f.cell(row=r, column=3).value,
            })

    summary["dashboard"] = {
        "active_scenario": {"cell": "Dashboard!B5", "value": clean_val(ws_dash["B5"].value)},
        "capex_mode": {"cell": "Dashboard!B6", "value": clean_val(ws_dash["B6"].value)},
        "kpis": dashboard_kpis,
    }

    # 4. Scenarios sheet (Scenarios!A6:G56) — Lease mode (saved, H4=2) vs Buy mode (computed, H4=1)
    ws_sc = wb_val["Scenarios"]
    ws_sc_f = wb_fml["Scenarios"]
    ws_cbuy = wb_val["CAPEX (Buy)"]
    ws_crent = wb_val["CAPEX (Rent)"]

    scenarios_rows = []
    for r in range(6, 57):
        label = clean_val(ws_sc.cell(row=r, column=1).value)
        unit = clean_val(ws_sc.cell(row=r, column=2).value)
        c_val = clean_val(ws_sc.cell(row=r, column=3).value)
        d_val = clean_val(ws_sc.cell(row=r, column=4).value)
        e_val = clean_val(ws_sc.cell(row=r, column=5).value)
        f_val = clean_val(ws_sc.cell(row=r, column=6).value)
        if label and (c_val is not None or d_val is not None):
            scenarios_rows.append({
                "row": r,
                "label": label,
                "unit": unit,
                "conservative_cell": f"Scenarios!C{r}",
                "conservative": c_val,
                "base_cell": f"Scenarios!D{r}",
                "base": d_val,
                "optimistic_cell": f"Scenarios!E{r}",
                "optimistic": e_val,
                "active_cell": f"Scenarios!F{r}",
                "active": f_val,
                "benchmark_note": clean_val(ws_sc.cell(row=r, column=7).value),
                "base_formula": ws_sc_f.cell(row=r, column=4).value,
            })

    # Compute exact Lease Mode (Inputs!H4 = 2, as saved) vs Buy Mode (Inputs!H4 = 1) across Conservative (col C), Base (col D), Optimistic (col E)
    # In Scenarios:
    # Row 30 Fixed overheads = SUM(Inputs!col67:col72) + IF(Inputs!$H$4=2, 'CAPEX (Rent)'!$H$47, 0)
    # Row 48 Hard costs = CHOOSE(Inputs!$H$4, 'CAPEX (Buy)'!$H$40, 'CAPEX (Rent)'!$H$43) * Inputs!col77
    # Row 49 Soft costs = CHOOSE(Inputs!$H$4, 'CAPEX (Buy)'!$H$41, 'CAPEX (Rent)'!$H$44) + Inputs!col79 * col26
    # Row 50 Contingency = Inputs!col78 * (col48 + col49)
    # Row 51 Lease (fit-out rent, prepayment, deposit) = col27*(Inputs!col31+Inputs!col29+Inputs!col30) + col28*Inputs!col31
    # Row 52 Working-capital reserve = Inputs!col80 * (col26 + col30 + col29)
    # Row 53 TOTAL INVESTMENT = SUM(col48:col52)
    lease_equip_pmt = ws_crent["H47"].value  # 18,483,149.583333332
    buy_hard_base = ws_cbuy["H40"].value     # 2,595,180,000
    buy_soft_base = ws_cbuy["H41"].value     # 677,100,000

    buy_vs_lease_scenarios = {}
    for scen_name, sc_col, in_col in [
        ("Conservative", 3, 5),  # Scenarios col C (3), Inputs col E (5)
        ("Base", 4, 6),          # Scenarios col D (4), Inputs col F (6)
        ("Optimistic", 5, 7),    # Scenarios col E (5), Inputs col G (7)
    ]:
        fx = ws_in.cell(row=10, column=in_col).value
        cit = ws_in.cell(row=16, column=in_col).value
        overrun = ws_in.cell(row=77, column=in_col).value
        contingency_rate = ws_in.cell(row=78, column=in_col).value
        preopen_pay_m = ws_in.cell(row=79, column=in_col).value
        wc_months = ws_in.cell(row=80, column=in_col).value
        equity_chf = ws_in.cell(row=82, column=in_col).value

        rev = ws_sc.cell(row=13, column=sc_col).value
        contrib = ws_sc.cell(row=23, column=sc_col).value
        contrib_pct = ws_sc.cell(row=24, column=sc_col).value
        staff = ws_sc.cell(row=26, column=sc_col).value
        rent = ws_sc.cell(row=27, column=sc_col).value
        sc_cam = ws_sc.cell(row=28, column=sc_col).value
        util = ws_sc.cell(row=29, column=sc_col).value

        # Saved Lease mode (H4 = 2)
        fixed_ovh_lease = ws_sc.cell(row=30, column=sc_col).value
        tot_fixed_lease = ws_sc.cell(row=31, column=sc_col).value
        ebitda_lease = ws_sc.cell(row=34, column=sc_col).value
        ebitda_margin_lease = ws_sc.cell(row=35, column=sc_col).value
        be_rev_mo_lease = ws_sc.cell(row=42, column=sc_col).value
        be_guests_day_lease = ws_sc.cell(row=45, column=sc_col).value
        hard_lease = ws_sc.cell(row=48, column=sc_col).value
        soft_lease = ws_sc.cell(row=49, column=sc_col).value
        cont_lease = ws_sc.cell(row=50, column=sc_col).value
        lease_preopen = ws_sc.cell(row=51, column=sc_col).value
        wc_lease = ws_sc.cell(row=52, column=sc_col).value
        tot_inv_lease_idr = ws_sc.cell(row=53, column=sc_col).value
        tot_inv_lease_chf = ws_sc.cell(row=54, column=sc_col).value
        gap_lease_chf = ws_sc.cell(row=55, column=sc_col).value
        payback_lease = ws_sc.cell(row=56, column=sc_col).value

        # Computed Buy mode (H4 = 1)
        fixed_ovh_buy = fixed_ovh_lease - lease_equip_pmt
        tot_fixed_buy = staff + rent + sc_cam + util + fixed_ovh_buy
        ebitda_buy = contrib - tot_fixed_buy
        ebitda_margin_buy = ebitda_buy / rev
        be_rev_mo_buy = tot_fixed_buy / contrib_pct
        guests_day = ws_sc.cell(row=6, column=sc_col).value
        be_guests_day_buy = guests_day * be_rev_mo_buy / rev

        hard_buy = buy_hard_base * overrun
        soft_buy = buy_soft_base + preopen_pay_m * staff
        cont_buy = contingency_rate * (hard_buy + soft_buy)
        wc_buy = wc_months * (staff + fixed_ovh_buy + util)
        tot_inv_buy_idr = hard_buy + soft_buy + cont_buy + lease_preopen + wc_buy
        tot_inv_buy_chf = tot_inv_buy_idr / fx
        gap_buy_chf = tot_inv_buy_chf - equity_chf
        payback_buy = (tot_inv_buy_idr / (ebitda_buy * 12 * (1 - cit))) if ebitda_buy > 0 else "never"

        buy_vs_lease_scenarios[scen_name] = {
            "fx_idr_per_chf": clean_val(fx),
            "guests_per_day_in_store": clean_val(guests_day),
            "avg_net_spend_per_guest_idr": clean_val(ws_sc.cell(row=8, column=sc_col).value),
            "monthly_revenue_idr": clean_val(rev),
            "annual_revenue_chf": clean_val(ws_sc.cell(row=14, column=sc_col).value),
            "contribution_margin_idr": clean_val(contrib),
            "contribution_margin_pct": clean_val(contrib_pct),
            "lease_mode_H4_eq_2_saved": {
                "fixed_overheads_incl_equip_idr": clean_val(fixed_ovh_lease),
                "total_fixed_costs_idr": clean_val(tot_fixed_lease),
                "monthly_ebitda_idr": clean_val(ebitda_lease),
                "ebitda_margin": clean_val(ebitda_margin_lease),
                "annual_ebitda_chf": clean_val(ws_sc.cell(row=36, column=sc_col).value),
                "breakeven_revenue_month_idr": clean_val(be_rev_mo_lease),
                "breakeven_guests_day": clean_val(be_guests_day_lease),
                "hard_costs_idr": clean_val(hard_lease),
                "soft_costs_idr": clean_val(soft_lease),
                "contingency_idr": clean_val(cont_lease),
                "lease_preopening_cash_idr": clean_val(lease_preopen),
                "working_capital_reserve_idr": clean_val(wc_lease),
                "total_investment_idr": clean_val(tot_inv_lease_idr),
                "total_investment_chf": clean_val(tot_inv_lease_chf),
                "funding_gap_vs_200k_chf": clean_val(gap_lease_chf),
                "simple_payback_after_tax_years": clean_val(payback_lease),
            },
            "buy_mode_H4_eq_1_computed": {
                "fixed_overheads_excl_equip_idr": clean_val(fixed_ovh_buy),
                "total_fixed_costs_idr": clean_val(tot_fixed_buy),
                "monthly_ebitda_idr": clean_val(ebitda_buy),
                "ebitda_margin": clean_val(ebitda_margin_buy),
                "annual_ebitda_chf": clean_val(ebitda_buy * 12 / fx),
                "breakeven_revenue_month_idr": clean_val(be_rev_mo_buy),
                "breakeven_guests_day": clean_val(be_guests_day_buy),
                "hard_costs_idr": clean_val(hard_buy),
                "soft_costs_idr": clean_val(soft_buy),
                "contingency_idr": clean_val(cont_buy),
                "lease_preopening_cash_idr": clean_val(lease_preopen),
                "working_capital_reserve_idr": clean_val(wc_buy),
                "total_investment_idr": clean_val(tot_inv_buy_idr),
                "total_investment_chf": clean_val(tot_inv_buy_chf),
                "funding_gap_vs_200k_chf": clean_val(gap_buy_chf),
                "simple_payback_after_tax_years": clean_val(payback_buy),
            },
        }

    summary["scenarios"] = {
        "rows": scenarios_rows,
        "lease_vs_buy_comparison": buy_vs_lease_scenarios,
    }

    # 5. Location & Rent_Data (Location!A7:U13, Rent_Data!A6:N39)
    ws_loc = wb_val["Location"]
    districts = []
    for r in range(8, 14):
        districts.append({
            "row": r,
            "district": clean_val(ws_loc.cell(row=r, column=1).value),
            "tier": clean_val(ws_loc.cell(row=r, column=2).value),
            "rent_p25_idr_sqm_mo": {"cell": f"Location!C{r}", "value": clean_val(ws_loc.cell(row=r, column=3).value)},
            "rent_median_idr_sqm_mo": {"cell": f"Location!D{r}", "value": clean_val(ws_loc.cell(row=r, column=4).value)},
            "rent_p75_idr_sqm_mo": {"cell": f"Location!E{r}", "value": clean_val(ws_loc.cell(row=r, column=5).value)},
            "sample_n": {"cell": f"Location!F{r}", "value": clean_val(ws_loc.cell(row=r, column=6).value)},
            "service_charge_idr_sqm_mo": {"cell": f"Location!G{r}", "value": clean_val(ws_loc.cell(row=r, column=7).value)},
            "demand_index": {"cell": f"Location!H{r}", "value": clean_val(ws_loc.cell(row=r, column=8).value)},
            "weighted_score": {"cell": f"Location!O{r}", "value": clean_val(ws_loc.cell(row=r, column=15).value)},
            "monthly_rent_plus_sc_idr": {"cell": f"Location!P{r}", "value": clean_val(ws_loc.cell(row=r, column=16).value)},
            "est_monthly_rev_idr": {"cell": f"Location!Q{r}", "value": clean_val(ws_loc.cell(row=r, column=17).value)},
            "occupancy_cost_pct": {"cell": f"Location!R{r}", "value": clean_val(ws_loc.cell(row=r, column=18).value)},
            "est_monthly_ebitda_idr": {"cell": f"Location!S{r}", "value": clean_val(ws_loc.cell(row=r, column=19).value)},
            "breakeven_guests_day": {"cell": f"Location!T{r}", "value": clean_val(ws_loc.cell(row=r, column=20).value)},
            "rent_check": {"cell": f"Location!U{r}", "value": clean_val(ws_loc.cell(row=r, column=21).value)},
        })
    summary["location_and_rent"] = {"districts": districts}

    # 6. Menu (Menu!A6:K49)
    ws_menu = wb_val["Menu"]
    menu_items = []
    for r in range(7, 27):
        cat = clean_val(ws_menu.cell(row=r, column=1).value)
        item = clean_val(ws_menu.cell(row=r, column=2).value)
        if item:
            menu_items.append({
                "row": r,
                "category": cat,
                "item": item,
                "price_net_idr": {"cell": f"Menu!C{r}", "value": clean_val(ws_menu.cell(row=r, column=3).value)},
                "food_cost_idr": {"cell": f"Menu!D{r}", "value": clean_val(ws_menu.cell(row=r, column=4).value)},
                "cost_pct": {"cell": f"Menu!E{r}", "value": clean_val(ws_menu.cell(row=r, column=5).value)},
                "contribution_idr": {"cell": f"Menu!F{r}", "value": clean_val(ws_menu.cell(row=r, column=6).value)},
                "units_per_100_guests": {"cell": f"Menu!G{r}", "value": clean_val(ws_menu.cell(row=r, column=7).value)},
                "rev_per_100_guests_idr": {"cell": f"Menu!H{r}", "value": clean_val(ws_menu.cell(row=r, column=8).value)},
                "cogs_per_100_guests_idr": {"cell": f"Menu!I{r}", "value": clean_val(ws_menu.cell(row=r, column=9).value)},
                "kasavana_smith_class": {"cell": f"Menu!J{r}", "value": clean_val(ws_menu.cell(row=r, column=10).value)},
            })

    menu_cats = []
    for r in range(33, 38):
        menu_cats.append({
            "row": r,
            "category": clean_val(ws_menu.cell(row=r, column=1).value),
            "revenue_share": {"cell": f"Menu!B{r}", "value": clean_val(ws_menu.cell(row=r, column=2).value)},
            "cost_pct": {"cell": f"Menu!C{r}", "value": clean_val(ws_menu.cell(row=r, column=3).value)},
        })

    menu_benchmarks = []
    for r in range(41, 50):
        menu_benchmarks.append({
            "row": r,
            "venue": clean_val(ws_menu.cell(row=r, column=1).value),
            "item": clean_val(ws_menu.cell(row=r, column=2).value),
            "price_idr": {"cell": f"Menu!C{r}", "value": clean_val(ws_menu.cell(row=r, column=3).value)},
            "channel": clean_val(ws_menu.cell(row=r, column=4).value),
            "area": clean_val(ws_menu.cell(row=r, column=5).value),
            "source_id": clean_val(ws_menu.cell(row=r, column=6).value),
        })

    summary["menu"] = {
        "items": menu_items,
        "blended_totals_row_27_to_30": {
            "units_per_100_guests": {"cell": "Menu!G27", "value": clean_val(ws_menu["G27"].value)},
            "rev_per_100_guests_idr": {"cell": "Menu!H27", "value": clean_val(ws_menu["H27"].value)},
            "cogs_per_100_guests_idr": {"cell": "Menu!I27", "value": clean_val(ws_menu["I27"].value)},
            "blended_fb_cost_pct": {"cell": "Menu!E29", "value": clean_val(ws_menu["E29"].value)},
            "implied_avg_net_spend_per_guest_idr": {"cell": "Menu!E30", "value": clean_val(ws_menu["E30"].value)},
            "revenue_model_blended_spend_idr": {"cell": "Menu!H30", "value": clean_val(ws_menu["H30"].value)},
        },
        "category_shares": menu_cats,
        "competitor_benchmarks": menu_benchmarks,
    }

    # 7. Revenue & Dayparts (Revenue!A5:K26)
    ws_rev = wb_val["Revenue"]
    dayparts = []
    for r in range(8, 13):
        dayparts.append({
            "row": r,
            "daypart": clean_val(ws_rev.cell(row=r, column=1).value),
            "hours": clean_val(ws_rev.cell(row=r, column=2).value),
            "guests_per_day": {"cell": f"Revenue!C{r}", "value": clean_val(ws_rev.cell(row=r, column=3).value)},
            "avg_net_spend_idr": {"cell": f"Revenue!D{r}", "value": clean_val(ws_rev.cell(row=r, column=4).value)},
            "daily_revenue_idr": {"cell": f"Revenue!E{r}", "value": clean_val(ws_rev.cell(row=r, column=5).value)},
            "takeaway_share": {"cell": f"Revenue!F{r}", "value": clean_val(ws_rev.cell(row=r, column=6).value)},
            "dine_in_guests": {"cell": f"Revenue!G{r}", "value": clean_val(ws_rev.cell(row=r, column=7).value)},
            "avg_dwell_min": {"cell": f"Revenue!H{r}", "value": clean_val(ws_rev.cell(row=r, column=8).value)},
            "max_covers_capacity": {"cell": f"Revenue!I{r}", "value": clean_val(ws_rev.cell(row=r, column=9).value)},
            "seat_utilisation_pct": {"cell": f"Revenue!J{r}", "value": clean_val(ws_rev.cell(row=r, column=10).value)},
            "revpash_idr": {"cell": f"Revenue!K{r}", "value": clean_val(ws_rev.cell(row=r, column=11).value)},
        })

    summary["revenue"] = {
        "seats": {"cell": "Revenue!C5", "value": clean_val(ws_rev["C5"].value)},
        "dayparts": dayparts,
        "delivery_row_13": {
            "orders_per_day": {"cell": "Revenue!C13", "value": clean_val(ws_rev["C13"].value)},
            "avg_order_value_idr": {"cell": "Revenue!D13", "value": clean_val(ws_rev["D13"].value)},
            "daily_delivery_revenue_idr": {"cell": "Revenue!E13", "value": clean_val(ws_rev["E13"].value)},
        },
        "total_fb_per_day_idr": {"cell": "Revenue!E14", "value": clean_val(ws_rev["E14"].value)},
        "blended_avg_net_spend_in_store_idr": {"cell": "Revenue!E16", "value": clean_val(ws_rev["E16"].value)},
        "mature_monthly_revenue": {
            "fb_in_store_idr": {"cell": "Revenue!E19", "value": clean_val(ws_rev["E19"].value)},
            "fb_delivery_idr": {"cell": "Revenue!E20", "value": clean_val(ws_rev["E20"].value)},
            "events_b2b_idr": {"cell": "Revenue!E21", "value": clean_val(ws_rev["E21"].value)},
            "retail_idr": {"cell": "Revenue!E22", "value": clean_val(ws_rev["E22"].value)},
            "total_monthly_idr": {"cell": "Revenue!E23", "value": clean_val(ws_rev["E23"].value)},
            "total_monthly_chf": {"cell": "Revenue!F23", "value": clean_val(ws_rev["F23"].value)},
            "annualised_idr": {"cell": "Revenue!E24", "value": clean_val(ws_rev["E24"].value)},
            "annualised_chf": {"cell": "Revenue!F24", "value": clean_val(ws_rev["F24"].value)},
        },
    }

    # 8. Staffing & OPEX (Staffing!A5:K18, OPEX!A5:D35)
    ws_st = wb_val["Staffing"]
    staff_roles = []
    for r in range(6, 15):
        staff_roles.append({
            "row": r,
            "role": clean_val(ws_st.cell(row=r, column=1).value),
            "headcount": {"cell": f"Staffing!B{r}", "value": clean_val(ws_st.cell(row=r, column=2).value)},
            "base_salary_idr": {"cell": f"Staffing!C{r}", "value": clean_val(ws_st.cell(row=r, column=3).value)},
            "gross_salaries_idr": {"cell": f"Staffing!F{r}", "value": clean_val(ws_st.cell(row=r, column=6).value)},
            "bpjs_employer_idr": {"cell": f"Staffing!G{r}", "value": clean_val(ws_st.cell(row=r, column=7).value)},
            "thr_accrual_idr": {"cell": f"Staffing!H{r}", "value": clean_val(ws_st.cell(row=r, column=8).value)},
            "benefits_idr": {"cell": f"Staffing!I{r}", "value": clean_val(ws_st.cell(row=r, column=9).value)},
            "total_monthly_cost_idr": {"cell": f"Staffing!J{r}", "value": clean_val(ws_st.cell(row=r, column=10).value)},
            "total_monthly_cost_chf": {"cell": f"Staffing!K{r}", "value": clean_val(ws_st.cell(row=r, column=11).value)},
        })

    ws_op = wb_val["OPEX"]
    opex_lines = []
    for r in range(7, 36):
        label = clean_val(ws_op.cell(row=r, column=1).value)
        val = clean_val(ws_op.cell(row=r, column=2).value)
        if label and val is not None:
            opex_lines.append({
                "row": r,
                "cost_line": label,
                "value": {"cell": f"OPEX!B{r}", "value": val},
                "driver": clean_val(ws_op.cell(row=r, column=3).value),
                "comment": clean_val(ws_op.cell(row=r, column=4).value),
            })

    summary["staffing_and_opex"] = {
        "staff_roles": staff_roles,
        "staff_total_row_15": {
            "total_headcount": {"cell": "Staffing!B15", "value": clean_val(ws_st["B15"].value)},
            "gross_salaries_idr": {"cell": "Staffing!F15", "value": clean_val(ws_st["F15"].value)},
            "bpjs_employer_idr": {"cell": "Staffing!G15", "value": clean_val(ws_st["G15"].value)},
            "thr_accrual_idr": {"cell": "Staffing!H15", "value": clean_val(ws_st["H15"].value)},
            "benefits_idr": {"cell": "Staffing!I15", "value": clean_val(ws_st["I15"].value)},
            "total_monthly_cost_idr": {"cell": "Staffing!J15", "value": clean_val(ws_st["J15"].value)},
            "total_monthly_cost_chf": {"cell": "Staffing!K15", "value": clean_val(ws_st["K15"].value)},
        },
        "opex_lines": opex_lines,
    }

    # 9. CAPEX (Buy), CAPEX (Rent), and Funding
    capex_buy_items = []
    for r in range(8, 38):
        item = clean_val(ws_cbuy.cell(row=r, column=2).value)
        if item:
            capex_buy_items.append({
                "row": r,
                "block": clean_val(ws_cbuy.cell(row=r, column=1).value),
                "item": item,
                "qty": clean_val(ws_cbuy.cell(row=r, column=3).value),
                "unit_cost_idr": clean_val(ws_cbuy.cell(row=r, column=4).value),
                "base_amount_idr": clean_val(ws_cbuy.cell(row=r, column=5).value),
                "final_amount_idr": {"cell": f"CAPEX (Buy)!H{r}", "value": clean_val(ws_cbuy.cell(row=r, column=8).value)},
                "final_amount_chf": {"cell": f"CAPEX (Buy)!I{r}", "value": clean_val(ws_cbuy.cell(row=r, column=9).value)},
                "share_of_total": {"cell": f"CAPEX (Buy)!J{r}", "value": clean_val(ws_cbuy.cell(row=r, column=10).value)},
            })

    capex_rent_items = []
    for r in range(8, 41):
        item = clean_val(ws_crent.cell(row=r, column=2).value)
        if item:
            capex_rent_items.append({
                "row": r,
                "block": clean_val(ws_crent.cell(row=r, column=1).value),
                "item": item,
                "base_amount_idr": clean_val(ws_crent.cell(row=r, column=5).value),
                "upfront_cash_idr": {"cell": f"CAPEX (Rent)!H{r}", "value": clean_val(ws_crent.cell(row=r, column=8).value)},
                "upfront_cash_chf": {"cell": f"CAPEX (Rent)!I{r}", "value": clean_val(ws_crent.cell(row=r, column=9).value)},
                "route_code": {"cell": f"CAPEX (Rent)!N{r}", "value": clean_val(ws_crent.cell(row=r, column=14).value)},
                "route_name": {"cell": f"CAPEX (Rent)!O{r}", "value": clean_val(ws_crent.cell(row=r, column=15).value)},
                "monthly_payment_idr": {"cell": f"CAPEX (Rent)!U{r}", "value": clean_val(ws_crent.cell(row=r, column=21).value)},
            })

    buy_vs_lease_summary_block = []
    for r in range(54, 62):
        metric = clean_val(ws_crent.cell(row=r, column=2).value)
        if metric:
            buy_vs_lease_summary_block.append({
                "row": r,
                "metric": metric,
                "buy_cell": f"CAPEX (Rent)!C{r}",
                "buy_val": clean_val(ws_crent.cell(row=r, column=3).value),
                "lease_cell": f"CAPEX (Rent)!D{r}",
                "lease_val": clean_val(ws_crent.cell(row=r, column=4).value),
                "delta_cell": f"CAPEX (Rent)!E{r}",
                "delta_val": clean_val(ws_crent.cell(row=r, column=5).value),
            })

    ws_fund = wb_val["Funding"]
    funding_rows = []
    for r in range(7, 26):
        item = clean_val(ws_fund.cell(row=r, column=1).value)
        if item and (ws_fund.cell(row=r, column=2).value is not None or ws_fund.cell(row=r, column=3).value is not None):
            funding_rows.append({
                "row": r,
                "item": item,
                "idr_or_val": {"cell": f"Funding!B{r}", "value": clean_val(ws_fund.cell(row=r, column=2).value)},
                "chf": {"cell": f"Funding!C{r}", "value": clean_val(ws_fund.cell(row=r, column=3).value)},
                "comment": clean_val(ws_fund.cell(row=r, column=4).value),
            })

    summary["capex_and_funding"] = {
        "capex_buy_items": capex_buy_items,
        "capex_rent_items": capex_rent_items,
        "buy_vs_lease_summary_block": buy_vs_lease_summary_block,
        "funding_rows": funding_rows,
    }

    # 10. Sensitivity & MonteCarlo (Sensitivity!A5:H64, MonteCarlo!A5:C34)
    ws_sens = wb_val["Sensitivity"]
    tornado_sorted = []
    for r in range(28, 35):
        tornado_sorted.append({
            "row": r,
            "driver": clean_val(ws_sens.cell(row=r, column=1).value),
            "delta_ebitda_minus_20pct_idr": {"cell": f"Sensitivity!B{r}", "value": clean_val(ws_sens.cell(row=r, column=2).value)},
            "delta_ebitda_plus_20pct_idr": {"cell": f"Sensitivity!C{r}", "value": clean_val(ws_sens.cell(row=r, column=3).value)},
            "swing_abs_idr": {"cell": f"Sensitivity!D{r}", "value": clean_val(ws_sens.cell(row=r, column=4).value)},
        })

    grid1_vol_spend = {
        "spend_factors": [clean_val(ws_sens.cell(row=39, column=c).value) for c in range(2, 9)],
        "rows": [
            {
                "volume_factor": clean_val(ws_sens.cell(row=r, column=1).value),
                "ebitda_idr": [clean_val(ws_sens.cell(row=r, column=c).value) for c in range(2, 9)],
            }
            for r in range(40, 47)
        ],
    }

    grid2_vol_rent = {
        "rent_factors": [clean_val(ws_sens.cell(row=49, column=c).value) for c in range(2, 9)],
        "rows": [
            {
                "volume_factor": clean_val(ws_sens.cell(row=r, column=1).value),
                "ebitda_idr": [clean_val(ws_sens.cell(row=r, column=c).value) for c in range(2, 9)],
            }
            for r in range(50, 57)
        ],
    }

    grid3_payback = {
        "ebitda_factors": [clean_val(ws_sens.cell(row=59, column=c).value) for c in range(2, 7)],
        "rows": [
            {
                "investment_factor": clean_val(ws_sens.cell(row=r, column=1).value),
                "payback_years": [clean_val(ws_sens.cell(row=r, column=c).value) for c in range(2, 7)],
            }
            for r in range(60, 65)
        ],
    }

    ws_mc = wb_val["MonteCarlo"]
    mc_drivers = []
    for r in range(16, 22):
        mc_drivers.append({
            "row": r,
            "driver": clean_val(ws_mc.cell(row=r, column=1).value),
            "min_factor": {"cell": f"MonteCarlo!B{r}", "value": clean_val(ws_mc.cell(row=r, column=2).value)},
            "mode_factor": {"cell": f"MonteCarlo!C{r}", "value": clean_val(ws_mc.cell(row=r, column=3).value)},
            "max_factor": {"cell": f"MonteCarlo!D{r}", "value": clean_val(ws_mc.cell(row=r, column=4).value)},
        })

    mc_results = []
    for r in range(24, 34):
        lbl = clean_val(ws_mc.cell(row=r, column=1).value)
        val = clean_val(ws_mc.cell(row=r, column=2).value)
        if lbl and val is not None:
            mc_results.append({
                "row": r,
                "metric": lbl,
                "cell": f"MonteCarlo!B{r}",
                "value": val,
            })

    summary["sensitivity_and_montecarlo"] = {
        "flex_pct": {"cell": "Sensitivity!E4", "value": clean_val(ws_sens["E4"].value)},
        "base_ebitda_month_idr": {"cell": "Sensitivity!B13", "value": clean_val(ws_sens["B13"].value)},
        "tornado_sorted": tornado_sorted,
        "grid1_volume_vs_spend": grid1_vol_spend,
        "grid2_volume_vs_rent": grid2_vol_rent,
        "grid3_payback": grid3_payback,
        "montecarlo_drivers": mc_drivers,
        "montecarlo_results": mc_results,
        "montecarlo_methodology_note": (
            "MonteCarlo!A40:K1039 runs 1,000 static draws of steady-state monthly EBITDA and simple after-tax "
            "steady-state payback around the Base scenario using triangular distributions (with volume & spend correlated "
            "via MonteCarlo!B33 = 1). It does NOT simulate 60-month cash flows or ramp-up working capital."
        ),
    }

    # 11. Sources sheet (Sources!A5:G108)
    ws_src = wb_val["Sources"]
    model_sources = []
    for r in range(5, 109):
        sid = clean_val(ws_src.cell(row=r, column=1).value)
        if sid:
            model_sources.append({
                "row": r,
                "source_id": sid,
                "category": clean_val(ws_src.cell(row=r, column=2).value),
                "parameter_supported": clean_val(ws_src.cell(row=r, column=3).value),
                "citation": clean_val(ws_src.cell(row=r, column=4).value),
                "url": clean_val(ws_src.cell(row=r, column=5).value),
                "date": str(ws_src.cell(row=r, column=6).value)[:10] if ws_src.cell(row=r, column=6).value else None,
                "notes": clean_val(ws_src.cell(row=r, column=7).value),
            })
    summary["model_sources_count"] = len(model_sources)
    summary["model_sources"] = model_sources

    json_out = OUT_DIR / "model_summary.json"
    with open(json_out, "w", encoding="utf-8") as f:
        json.dump(summary, f, indent=2, ensure_ascii=False)

    # Write Markdown Tables for Section 8
    md_lines = [
        "# CORSA Jakarta Financial Model v0.6 — Extracted Figures & Cell Citations",
        "",
        f"- **Source file:** `{summary['model_file']}`",
        f"- **Sheets ({summary['sheet_count']}):** {', '.join(summary['sheet_names'])}",
        f"- **Active Scenario Switch:** `Inputs!D4 = {summary['switches_and_inputs']['active_scenario_code']['value']}` (`Inputs!E4 = \"{summary['switches_and_inputs']['active_scenario_name']['value']}\"`)",
        f"- **Active CAPEX Financing Switch:** `Inputs!H4 = {summary['switches_and_inputs']['capex_financing_mode_code']['value']}` (`Inputs!I4 = \"{summary['switches_and_inputs']['capex_financing_mode_name']['value']}\"`)",
        "",
        "## 1. Structural Completeness & Horizon Audit (`PnL_Monthly`, `Annual`, `Dashboard`, `Research_Tracker`)",
        "",
        "- **`PnL_Monthly` horizon truncation:** Only **13 of 60 months** are populated (`PnL_Monthly!C4:O53`, Months 1–13 = Jul 2027 – Jul 2028). Columns `P` through `BJ` (Months 14–60) are completely blank (`PnL_Monthly!P4:BJ53`).",
        "- **`Annual` distortion (`Annual!A4:H32`):**",
        "  - Year 1 (`Annual!C5:C19`) sums operating year 1 (Months 1–12): Revenue = `Rp 7,295,392,080` (`CHF 337,750`, `Annual!C21`), EBITDA = `Rp 730,264,586` (`CHF 33,809`, `Annual!C22`), Net Income = `Rp 293,511,415` (`Annual!C16`), Project FCF = `Rp 1,100,769,940` (`Annual!C17`), Closing Cash = `Rp 1,914,702,149` (`Annual!C19`).",
        "  - Year 2 (`Annual!D5:D19`) sums operating year 2 (Months 13–24), where only Month 13 (`PnL_Monthly!O`) is populated: Revenue = `Rp 765,055,200` (`CHF 35,419`), EBITDA = `Rp 136,829,613` (`CHF 6,335`), Project FCF = `Rp 113,216,510`.",
        "  - Years 3, 4, and 5 (`Annual!E5:G19`) evaluate to **`0`** across all revenue, cost, EBITDA, and operating cash-flow lines (only `Annual!G29` has `Rp 75,548,438` from the refundable security deposit).",
        "- **Invalidated `Dashboard` multi-year metrics:**",
        "  - `Dashboard!B17:C17` (Year-3 EBITDA) = `Rp 0` / `CHF 0` (references `Annual!E11`).",
        "  - `Dashboard!B19:C19` & `Annual!C31:D31` (NPV of 5-year project cash flows at 15% discount rate) = `Rp -3,185,886,005` / `CHF -147,495`.",
        "  - `Dashboard!B25` (Payback month, full model) = `> 60 months`.",
        "  - `Dashboard!B26` & `Annual!C30` (Project IRR, 5 yrs) = `-46.9%`.",
        "  - `Dashboard!B18:C18` (Lowest cash balance over 60 months) = `Rp 794,813,727` / `CHF 36,797` (Month 1 cash balance in `PnL_Monthly!C48`, ignoring unpopulated Months 14–60).",
        "- **`Research_Tracker` status (`Research_Tracker!A5:F17`):** 13 rows — **5 `Placeholder`**, **4 `Desk`**, **4 `Model ready`**.",
        "",
        "## 2. Side-by-Side Scenario Table (`Scenarios!A6:G56`, Saved Lease Mode `Inputs!H4 = 2`)",
        "",
        "| Row | Metric | Unit | Conservative (`C`) | Base (`D`) | Optimistic (`E`) | Cell Reference |",
        "|---:|---|---|---:|---:|---:|---|",
    ]

    for row in scenarios_rows:
        r = row["row"]
        unit = row["unit"] or ""
        c_v, b_v, o_v = row["conservative"], row["base"], row["optimistic"]
        if unit == "IDR":
            cs, bs, os_ = fmt_idr(c_v), fmt_idr(b_v), fmt_idr(o_v)
        elif unit == "CHF":
            cs, bs, os_ = fmt_chf(c_v), fmt_chf(b_v), fmt_chf(o_v)
        elif unit == "%":
            cs, bs, os_ = fmt_pct(c_v), fmt_pct(b_v), fmt_pct(o_v)
        elif isinstance(b_v, float):
            cs = f"{c_v:,.2f}" if isinstance(c_v, (int, float)) else str(c_v)
            bs = f"{b_v:,.2f}"
            os_ = f"{o_v:,.2f}" if isinstance(o_v, (int, float)) else str(o_v)
        else:
            cs, bs, os_ = str(c_v), str(b_v), str(o_v)
        md_lines.append(f"| {r} | {row['label']} | {unit} | {cs} | {bs} | {os_} | `Scenarios!C{r}:E{r}` |")

    md_lines.extend([
        "",
        "## 3. Lease Mode (`Inputs!H4 = 2`, Saved) vs Buy Mode (`Inputs!H4 = 1`, Computed) Comparison",
        "",
        "| Metric | Conservative (Lease / Buy) | Base (Lease / Buy) | Optimistic (Lease / Buy) | Citations |",
        "|---|---:|---:|---:|---|",
    ])
    c_c = buy_vs_lease_scenarios["Conservative"]
    b_c = buy_vs_lease_scenarios["Base"]
    o_c = buy_vs_lease_scenarios["Optimistic"]
    md_lines.extend([
        f"| Steady-State Monthly Revenue | {fmt_idr(c_c['monthly_revenue_idr'])} | {fmt_idr(b_c['monthly_revenue_idr'])} | {fmt_idr(o_c['monthly_revenue_idr'])} | `Scenarios!C13:E13` |",
        f"| Fixed Overheads / mo (incl. equip. lease in Lease mode) | {fmt_idr(c_c['lease_mode_H4_eq_2_saved']['fixed_overheads_incl_equip_idr'])} / {fmt_idr(c_c['buy_mode_H4_eq_1_computed']['fixed_overheads_excl_equip_idr'])} | {fmt_idr(b_c['lease_mode_H4_eq_2_saved']['fixed_overheads_incl_equip_idr'])} / {fmt_idr(b_c['buy_mode_H4_eq_1_computed']['fixed_overheads_excl_equip_idr'])} | {fmt_idr(o_c['lease_mode_H4_eq_2_saved']['fixed_overheads_incl_equip_idr'])} / {fmt_idr(o_c['buy_mode_H4_eq_1_computed']['fixed_overheads_excl_equip_idr'])} | `Scenarios!C30:E30`, `CAPEX (Rent)!H47` |",
        f"| Total Fixed Costs / mo | {fmt_idr(c_c['lease_mode_H4_eq_2_saved']['total_fixed_costs_idr'])} / {fmt_idr(c_c['buy_mode_H4_eq_1_computed']['total_fixed_costs_idr'])} | {fmt_idr(b_c['lease_mode_H4_eq_2_saved']['total_fixed_costs_idr'])} / {fmt_idr(b_c['buy_mode_H4_eq_1_computed']['total_fixed_costs_idr'])} | {fmt_idr(o_c['lease_mode_H4_eq_2_saved']['total_fixed_costs_idr'])} / {fmt_idr(o_c['buy_mode_H4_eq_1_computed']['total_fixed_costs_idr'])} | `Scenarios!C31:E31` |",
        f"| Steady-State Monthly EBITDA (IDR) | {fmt_idr(c_c['lease_mode_H4_eq_2_saved']['monthly_ebitda_idr'])} / {fmt_idr(c_c['buy_mode_H4_eq_1_computed']['monthly_ebitda_idr'])} | {fmt_idr(b_c['lease_mode_H4_eq_2_saved']['monthly_ebitda_idr'])} / {fmt_idr(b_c['buy_mode_H4_eq_1_computed']['monthly_ebitda_idr'])} | {fmt_idr(o_c['lease_mode_H4_eq_2_saved']['monthly_ebitda_idr'])} / {fmt_idr(o_c['buy_mode_H4_eq_1_computed']['monthly_ebitda_idr'])} | `Scenarios!C34:E34` |",
        f"| Steady-State EBITDA Margin (%) | {fmt_pct(c_c['lease_mode_H4_eq_2_saved']['ebitda_margin'])} / {fmt_pct(c_c['buy_mode_H4_eq_1_computed']['ebitda_margin'])} | {fmt_pct(b_c['lease_mode_H4_eq_2_saved']['ebitda_margin'])} / {fmt_pct(b_c['buy_mode_H4_eq_1_computed']['ebitda_margin'])} | {fmt_pct(o_c['lease_mode_H4_eq_2_saved']['ebitda_margin'])} / {fmt_pct(o_c['buy_mode_H4_eq_1_computed']['ebitda_margin'])} | `Scenarios!C35:E35` |",
        f"| Break-Even Guests / Day | {c_c['lease_mode_H4_eq_2_saved']['breakeven_guests_day']:.1f} / {c_c['buy_mode_H4_eq_1_computed']['breakeven_guests_day']:.1f} | {b_c['lease_mode_H4_eq_2_saved']['breakeven_guests_day']:.1f} / {b_c['buy_mode_H4_eq_1_computed']['breakeven_guests_day']:.1f} | {o_c['lease_mode_H4_eq_2_saved']['breakeven_guests_day']:.1f} / {o_c['buy_mode_H4_eq_1_computed']['breakeven_guests_day']:.1f} | `Scenarios!C45:E45` |",
        f"| Total Investment Requirement (IDR) | {fmt_idr(c_c['lease_mode_H4_eq_2_saved']['total_investment_idr'])} / {fmt_idr(c_c['buy_mode_H4_eq_1_computed']['total_investment_idr'])} | {fmt_idr(b_c['lease_mode_H4_eq_2_saved']['total_investment_idr'])} / {fmt_idr(b_c['buy_mode_H4_eq_1_computed']['total_investment_idr'])} | {fmt_idr(o_c['lease_mode_H4_eq_2_saved']['total_investment_idr'])} / {fmt_idr(o_c['buy_mode_H4_eq_1_computed']['total_investment_idr'])} | `Scenarios!C53:E53`, `CAPEX (Rent)!C54:D54` |",
        f"| Total Investment Requirement (CHF) | {fmt_chf(c_c['lease_mode_H4_eq_2_saved']['total_investment_chf'])} / {fmt_chf(c_c['buy_mode_H4_eq_1_computed']['total_investment_chf'])} | {fmt_chf(b_c['lease_mode_H4_eq_2_saved']['total_investment_chf'])} / {fmt_chf(b_c['buy_mode_H4_eq_1_computed']['total_investment_chf'])} | {fmt_chf(o_c['lease_mode_H4_eq_2_saved']['total_investment_chf'])} / {fmt_chf(o_c['buy_mode_H4_eq_1_computed']['total_investment_chf'])} | `Scenarios!C54:E54`, `CAPEX (Rent)!C55:D55` |",
        f"| Gap (+) / Surplus (-) vs CHF 200k Equity | {fmt_chf(c_c['lease_mode_H4_eq_2_saved']['funding_gap_vs_200k_chf'])} / {fmt_chf(c_c['buy_mode_H4_eq_1_computed']['funding_gap_vs_200k_chf'])} | {fmt_chf(b_c['lease_mode_H4_eq_2_saved']['funding_gap_vs_200k_chf'])} / {fmt_chf(b_c['buy_mode_H4_eq_1_computed']['funding_gap_vs_200k_chf'])} | {fmt_chf(o_c['lease_mode_H4_eq_2_saved']['funding_gap_vs_200k_chf'])} / {fmt_chf(o_c['buy_mode_H4_eq_1_computed']['funding_gap_vs_200k_chf'])} | `Scenarios!C55:E55`, `CAPEX (Rent)!C56:D56` |",
        f"| Simple Steady-State Payback After Tax (Yrs) | never / never | {b_c['lease_mode_H4_eq_2_saved']['simple_payback_after_tax_years']:.2f} / {b_c['buy_mode_H4_eq_1_computed']['simple_payback_after_tax_years']:.2f} | {o_c['lease_mode_H4_eq_2_saved']['simple_payback_after_tax_years']:.2f} / {o_c['buy_mode_H4_eq_1_computed']['simple_payback_after_tax_years']:.2f} | `Scenarios!C56:E56` |",
    ])

    md_lines.extend([
        "",
        "## 4. Menu Engineering Table (`Menu!A6:J37`)",
        "",
        "| Row | Category | Item | Net Price (IDR) | Food Cost (IDR) | Cost % | Unit Contribution (IDR) | Units / 100 Guests | Class | Cell Ref |",
        "|---:|---|---|---:|---:|---:|---:|---:|---|---|",
    ])
    for m in menu_items:
        r = m["row"]
        md_lines.append(
            f"| {r} | {m['category']} | {m['item']} | {fmt_idr(m['price_net_idr']['value'])} | "
            f"{fmt_idr(m['food_cost_idr']['value'])} | {fmt_pct(m['cost_pct']['value'])} | "
            f"{fmt_idr(m['contribution_idr']['value'])} | {m['units_per_100_guests']['value']} | "
            f"{m['kasavana_smith_class']['value']} | `Menu!A{r}:J{r}` |"
        )

    md_lines.extend([
        "",
        "## 5. Sensitivity Tornado (`Sensitivity!A27:D34`) & Monte Carlo (`MonteCarlo!A24:B33`)",
        "",
        "| Rank | Driver (±20% flex around Base) | Δ EBITDA @ -20% (IDR) | Δ EBITDA @ +20% (IDR) | Total Swing (IDR) | Cell Ref |",
        "|---:|---|---:|---:|---:|---|",
    ])
    for idx, t in enumerate(tornado_sorted, 1):
        r = t["row"]
        md_lines.append(
            f"| {idx} | {t['driver']} | {fmt_idr(t['delta_ebitda_minus_20pct_idr']['value'])} | "
            f"{fmt_idr(t['delta_ebitda_plus_20pct_idr']['value'])} | {fmt_idr(t['swing_abs_idr']['value'])} | `Sensitivity!A{r}:D{r}` |"
        )

    md_lines.extend([
        "",
        "### Monte Carlo Summary (1,000 Steady-State Draws, `MonteCarlo!A24:B33`)",
        "",
        "| Metric | Value | Cell Ref |",
        "|---|---:|---|",
    ])
    for mr in mc_results:
        v = mr["value"]
        if "EBITDA (" in mr["metric"] or "monthly EBITDA" in mr["metric"]:
            vs = fmt_idr(v)
        elif "Probability" in mr["metric"]:
            vs = fmt_pct(v)
        elif isinstance(v, float):
            vs = f"{v:.2f}"
        else:
            vs = str(v)
        md_lines.append(f"| {mr['metric']} | {vs} | `{mr['cell']}` |")

    md_out = OUT_DIR / "model_figures_table.md"
    with open(md_out, "w", encoding="utf-8") as f:
        f.write("\n".join(md_lines) + "\n")

    print(f"[extract_model] Wrote {json_out} and {md_out}")
    return summary


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Extract CORSA financial model v0.6 figures")
    parser.add_argument("--model-path", type=Path, default=DEFAULT_MODEL, help="Path to financial model .xlsx")
    args = parser.parse_args()
    extract_model(args.model_path)
